"""
Inventory Manager - backend
Flask web server + Excel file (inventory.xlsx) as the database.

Run:        python app.py
Phone cam:  python app.py --https     (camera on a phone needs https)
"""
import os
import shutil
import socket
import sys
import threading
from datetime import datetime

from flask import Flask, jsonify, request, send_from_directory
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = os.path.dirname(os.path.abspath(__file__))
# Optional: set these environment variables to keep the data / backups in another folder
# (for example a Google Drive or OneDrive folder that syncs to the cloud).
XLSX = os.environ.get("INVENTORY_XLSX") or os.path.join(BASE, "inventory.xlsx")
BACKUPS = os.environ.get("INVENTORY_BACKUP_DIR") or os.path.join(BASE, "backups")
LOCK = threading.Lock()

ITEM_COLS = ["Code", "Item name", "Stock", "Cost price", "Selling price", "Low stock alert"]
LOG_COLS = ["ID", "Bill", "Date & time", "Type", "Code", "Item name", "Qty", "Stock after",
            "Cost price", "Selling price", "Payment", "Customer", "Phone", "Paid on"]
UNPAID_COLS = ["Customer", "Phone", "Item", "Qty", "Amount", "Date & time"]

app = Flask(__name__, static_folder=os.path.join(BASE, "static"))


# ---------- helpers ----------
def ms_to_dt(ms):
    return datetime.fromtimestamp((ms or 0) / 1000).replace(microsecond=0)


def dt_to_ms(v):
    return int(v.timestamp() * 1000) if isinstance(v, datetime) else 0


def num(v, default=0):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return default
    return int(f) if f == int(f) else f


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def style_sheet(ws, widths):
    head = PatternFill("solid", fgColor="2563EB")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


# ---------- Excel read / write ----------
def read_data():
    data = {"items": {}, "log": []}
    if not os.path.exists(XLSX):
        restore_from_backup()
    if not os.path.exists(XLSX):
        return data
    wb = load_workbook(XLSX, data_only=True)
    if "Items" in wb.sheetnames:
        for r in wb["Items"].iter_rows(min_row=2, values_only=True):
            if not r or r[0] in (None, ""):
                continue
            r = list(r) + [None] * (6 - len(r))
            data["items"][txt(r[0])] = {
                "name": txt(r[1]), "qty": num(r[2]), "cost": num(r[3]),
                "price": num(r[4]), "min": num(r[5], 5),
            }
    if "History" in wb.sheetnames:
        for r in wb["History"].iter_rows(min_row=2, values_only=True):
            if not r or r[0] in (None, ""):
                continue
            r = list(r) + [None] * (len(LOG_COLS) - len(r))
            e = {"id": txt(r[0]), "bill": txt(r[1]), "t": dt_to_ms(r[2]), "type": txt(r[3]),
                 "code": txt(r[4]), "name": txt(r[5]), "n": num(r[6]), "after": num(r[7])}
            if e["type"] == "OUT":
                e["cost"], e["price"] = num(r[8]), num(r[9])
                if r[10] in ("Unpaid", "Paid later"):
                    e["credit"] = True
                    e["paid"] = r[10] == "Paid later"
                    e["cust"], e["phone"] = txt(r[11]), txt(r[12])
                    if e["paid"]:
                        e["paidAt"] = dt_to_ms(r[13])
            data["log"].append(e)
    data["log"].sort(key=lambda e: e["t"], reverse=True)
    return data


def build_workbook(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(ITEM_COLS)
    items = sorted(data["items"].items(), key=lambda kv: str(kv[1].get("name", "")).lower())
    for code, it in items:
        ws.append([str(code), it.get("name", ""), it.get("qty", 0), it.get("cost", 0),
                   it.get("price", 0), it.get("min", 5)])
        ws.cell(row=ws.max_row, column=1).number_format = "@"
    style_sheet(ws, [20, 32, 10, 12, 14, 16])

    wl = wb.create_sheet("History")
    wl.append(LOG_COLS)
    for e in sorted(data["log"], key=lambda x: x.get("t", 0)):
        out = e.get("type") == "OUT"
        pay = ""
        if out:
            pay = "Paid" if not e.get("credit") else ("Paid later" if e.get("paid") else "Unpaid")
        wl.append([e.get("id"), e.get("bill", ""), ms_to_dt(e.get("t")), e.get("type"),
                   str(e.get("code", "")), e.get("name"), e.get("n"), e.get("after"),
                   e.get("cost") if out else None, e.get("price") if out else None, pay,
                   e.get("cust") if e.get("credit") else None,
                   str(e.get("phone")) if e.get("phone") else None,
                   ms_to_dt(e["paidAt"]) if e.get("paidAt") else None])
        row = wl.max_row
        wl.cell(row=row, column=3).number_format = "dd-mmm-yyyy hh:mm"
        wl.cell(row=row, column=14).number_format = "dd-mmm-yyyy hh:mm"
    style_sheet(wl, [14, 14, 20, 7, 18, 30, 7, 12, 11, 13, 12, 22, 14, 20])

    # read-only convenience sheet (rebuilt on every save, ignored on load)
    wu = wb.create_sheet("Unpaid")
    wu.append(UNPAID_COLS)
    pend = [e for e in data["log"] if e.get("type") == "OUT" and e.get("credit") and not e.get("paid")]
    for e in sorted(pend, key=lambda x: (str(x.get("cust", "")), x.get("t", 0))):
        wu.append([e.get("cust"), e.get("phone") or None, e.get("name"), e.get("n"),
                   (e.get("n") or 0) * (e.get("price") or 0), ms_to_dt(e.get("t"))])
        wu.cell(row=wu.max_row, column=6).number_format = "dd-mmm-yyyy hh:mm"
    style_sheet(wu, [24, 14, 30, 7, 12, 20])
    return wb


def backup(name=None):
    if not os.path.exists(XLSX):
        return
    os.makedirs(BACKUPS, exist_ok=True)
    dst = os.path.join(BACKUPS, name or "inventory-%s.xlsx" % datetime.now().strftime("%Y-%m-%d"))
    if name or not os.path.exists(dst):  # daily backup: first save of the day only
        shutil.copy2(XLSX, dst)


def sync_latest():
    """Keep a copy of the newest data in the backup folder (used for restore)."""
    try:
        os.makedirs(BACKUPS, exist_ok=True)
        shutil.copy2(XLSX, os.path.join(BACKUPS, "inventory-latest.xlsx"))
    except OSError:
        pass  # e.g. a cloud-sync lock: never fail the real save because of this


def restore_from_backup():
    """If inventory.xlsx is missing, load the newest valid Excel file from the backup folder."""
    if os.path.exists(XLSX) or not os.path.isdir(BACKUPS):
        return None
    files = [os.path.join(BACKUPS, f) for f in os.listdir(BACKUPS)
             if f.lower().endswith(".xlsx") and not f.startswith("~$") and "before-reset" not in f]
    files.sort(key=os.path.getmtime, reverse=True)
    for f in files:
        try:
            wb = load_workbook(f, read_only=True)
            ok = "Items" in wb.sheetnames
            wb.close()
        except Exception:
            continue
        if ok:
            shutil.copy2(f, XLSX)
            return f
    return None


def save_data(data):
    tmp = XLSX + ".tmp"
    try:
        build_workbook(data).save(tmp)
        backup()
        os.replace(tmp, XLSX)
        sync_latest()
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


# ---------- routes ----------
@app.get("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.get("/api/data")
def get_data():
    with LOCK:
        try:
            return jsonify(read_data())
        except Exception as ex:  # e.g. corrupt file
            return jsonify({"error": "Could not read inventory.xlsx: %s" % ex}), 500


@app.post("/api/data")
def post_data():
    d = request.get_json(silent=True)
    if not isinstance(d, dict) or not isinstance(d.get("items"), dict) or not isinstance(d.get("log"), list):
        return jsonify({"error": "Bad data"}), 400
    with LOCK:
        try:
            save_data({"items": d["items"], "log": d["log"]})
        except PermissionError:
            return jsonify({"error": "inventory.xlsx is open in Excel. Close it and try again."}), 409
        except Exception as ex:
            return jsonify({"error": str(ex)}), 500
    return jsonify({"ok": True})


@app.post("/api/reset")
def reset():
    with LOCK:
        try:
            backup("inventory-before-reset-%s.xlsx" % datetime.now().strftime("%Y%m%d-%H%M%S"))
            save_data({"items": {}, "log": []})
        except PermissionError:
            return jsonify({"error": "inventory.xlsx is open in Excel. Close it and try again."}), 409
    return jsonify({"ok": True})


def local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except OSError:
        return "127.0.0.1"


if __name__ == "__main__":
    use_https = "--https" in sys.argv
    scheme = "https" if use_https else "http"
    restored = restore_from_backup()
    if not os.path.exists(XLSX):
        save_data({"items": {}, "log": []})
    print("\n  Inventory Manager is running")
    if restored:
        print("  Data restored from backup: %s" % restored)
    print("  On this computer : %s://localhost:5000" % scheme)
    print("  On your phone    : %s://%s:5000   (same Wi-Fi)" % (scheme, local_ip()))
    print("  Data file        : %s" % XLSX)
    print("  Backups folder   : %s\n" % BACKUPS)
    app.run(host="0.0.0.0", port=5000, ssl_context="adhoc" if use_https else None)
